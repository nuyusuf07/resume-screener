"""
parser.py

Turns an uploaded resume file (PDF, DOCX, or TXT) into plain text so it can
be scored. Works with both real file paths and Streamlit's in-memory
UploadedFile objects.
"""

import io
import re

import pdfplumber
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph


def _read_pdf(file_obj) -> str:
    text_chunks = []
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def _extract_table_text(table) -> list:
    parts = []
    for row in table.rows:
        for cell in row.cells:
            cell_text = cell.text.strip()
            if cell_text:
                parts.append(cell_text)
            # CVs built with layout tables sometimes nest a table inside a cell
            for nested in cell.tables:
                parts.extend(_extract_table_text(nested))
    return parts


def _iter_block_items(document):
    """Yield paragraphs and tables in their actual document order."""
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _read_docx(file_obj) -> str:
    document = DocxDocument(file_obj)
    parts = []

    # `document.paragraphs` followed by `document.tables` changes the reading
    # order.  ATS checks need the same sequence a recruiter sees, so walk the
    # body XML and emit paragraphs/tables where they occur.
    for block in _iter_block_items(document):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text)
        else:
            parts.extend(_extract_table_text(block))

    return "\n".join(parts)


def _read_txt(file_obj) -> str:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        return raw.decode("utf-8", errors="ignore")
    return raw


def extract_text(uploaded_file) -> str:
    """
    uploaded_file: a Streamlit UploadedFile (has .name) or a plain file path (str).
    Returns cleaned plain text.
    """
    if isinstance(uploaded_file, str):
        name = uploaded_file
        file_obj = open(uploaded_file, "rb")
    else:
        name = uploaded_file.name
        file_obj = uploaded_file

    ext = name.lower().rsplit(".", 1)[-1]

    if ext == "pdf":
        text = _read_pdf(file_obj)
    elif ext == "docx":
        text = _read_docx(file_obj)
    elif ext == "txt":
        text = _read_txt(file_obj)
    else:
        raise ValueError(f"Unsupported file type: .{ext}. Please upload a PDF, DOCX, or TXT file.")

    return clean_text(text)


def clean_text(text: str) -> str:
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


_JOB_BOARD_NOISE = [
    r"view jobs(?: in .*)?",
    r"save",
    r"email",
    r"type your email(?: here)?\.*",
    r"sign in",
    r"log in",
    r"apply now",
    r"share(?: this job)?",
    r"job alert(?:s)?",
]


def clean_job_description(text: str) -> tuple[str, list[str]]:
    """
    Remove common job-board controls without deleting substantive JD content.

    Returns `(cleaned_text, removed_lines)` so the UI can make the cleanup
    visible.  Only whole-line boilerplate is removed; a real requirement that
    happens to contain words such as "email" remains untouched.
    """
    cleaned = clean_text(text)
    kept = []
    removed = []
    patterns = [re.compile(rf"^\s*{p}\s*$", re.IGNORECASE) for p in _JOB_BOARD_NOISE]

    for raw_line in cleaned.splitlines():
        line = raw_line.strip()
        if line and any(pattern.match(line) for pattern in patterns):
            removed.append(line)
        else:
            kept.append(raw_line)

    return clean_text("\n".join(kept)), removed
